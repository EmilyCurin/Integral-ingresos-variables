"""Flexible offline import. All extracted transactions remain reviewable proposals."""
import csv, io, os, re, unicodedata
from datetime import date, datetime
from pathlib import Path
from .core import date_normalize, parse_amount

MAX_BYTES = 12 * 1024 * 1024
ALLOWED = {'.csv', '.txt', '.xlsx', '.pdf', '.png', '.jpg', '.jpeg', '.webp'}
DATE_RE = re.compile(r'\b(?:\d{4}[-/]\d{1,2}[-/]\d{1,2}|\d{1,2}[-/]\d{1,2}[-/]\d{2,4})\b')
MONEY_RE = re.compile(r'(?<![\w])[-+]?(?:Q\s*)?\d{1,3}(?:[, .]\d{3})*[.,]\d{2}(?!\d)|(?<![\w])[-+]?(?:Q\s*)?\d+(?!\d)')

def norm(x):
    return ''.join(c for c in unicodedata.normalize('NFKD', str(x or '').lower()) if not unicodedata.combining(c)).strip()

def classify(desc, incoming=True):
    s=norm(desc)
    if not incoming: return 'gasto', 'alta'
    if any(x in s for x in ('reembolso', 'devolucion', 'reintegro')): return 'reembolso','media'
    if any(x in s for x in ('prestamo', 'credito personal')): return 'prestamo','media'
    # Transfers from an unspecified account are *not* automatically own-account transfers.
    if ('transf de cta' in s or 'transferencia propia' in s or 'entre mis cuentas' in s): return 'pendiente','baja'
    if any(x in s for x in ('clases','clase ','tutoria','tutoras','ajedrez','curso','mensualidad','inscripc','maqueta','servicio profesional','honorarios')):
        return 'variable','media'
    return 'pendiente','baja'

def proposal(day, amount, description='', source='', incoming=True, line=None):
    category,confidence=classify(description,incoming)
    return {'date':day, 'amount_cents':abs(amount),'description':str(description).strip()[:250],
            'source':str(source).strip()[:100], 'category':category,'note':'','confidence':confidence,'origin_line':line}

def is_date(v):
    if isinstance(v,datetime): return v.date().isoformat()
    if isinstance(v,date): return v.isoformat()
    try: return date_normalize(str(v).strip())
    except (ValueError,TypeError): return None

def is_money(v):
    if v is None or isinstance(v,bool): return None
    try:
        if isinstance(v,(int,float)):
            if not v: return None
            from decimal import Decimal
            return int(Decimal(str(v))*100)
        if str(v).strip() in ('','-','—'):return None
        return parse_amount(v)
    except (TypeError,ValueError):return None

def structured_rows(rows):
    """Accept any header row position and flexible names, including separate debit/credit."""
    if not rows:return []
    aliases={
        'date':['fecha','date','dia','día','fecha transaccion','fecha de movimiento'],
        'amount':['monto','cantidad','amount','importe','ingreso','valor','total'],
        'credit':['credito','crédito','abono','deposito','depósito','entrada','haber'],
        'debit':['debito','débito','cargo','retiro','salida','debe'],
        'description':['descripcion','descripción','description','concepto','detalle','movimiento','observacion','referencia'],
        'source':['fuente','source','cliente','origen','tercero','remitente']}
    header_at=None;cols={}
    for idx,row in enumerate(rows[:35]):
        names=[norm(x) for x in row]
        candidate={k:next((j for j,n in enumerate(names) if n in [norm(a) for a in options]),None) for k,options in aliases.items()}
        if candidate['date'] is not None and any(candidate[k] is not None for k in ('amount','credit','debit')):
            header_at=idx; cols=candidate;break
    if header_at is None:
        # Headerless notes with ISO/day-month-year date and numerical value next to it.
        for idx,row in enumerate(rows[:20]):
            d=next((j for j,v in enumerate(row) if is_date(v)),None)
            if d is None:continue
            m=next((j for j,v in enumerate(row) if j!=d and is_money(v) is not None),None)
            if m is not None:cols={'date':d,'amount':m,'credit':None,'debit':None,'description':None,'source':None};header_at=idx-1;break
    if header_at is None:raise ValueError('No identifiqué las fechas y los importes. Revisa si el archivo contiene movimientos; puedes ingresar o corregir los datos manualmente.')
    out=[]
    for line,row in enumerate(rows[header_at+1:],header_at+2):
        def get(k):
            j=cols.get(k);return row[j] if j is not None and j<len(row) else None
        dt=is_date(get('date'))
        if not dt:continue
        c=is_money(get('credit'));d=is_money(get('debit')); a=is_money(get('amount'))
        if c is not None and c>0:amt,incoming=c,True
        elif d is not None and d>0:amt,incoming=d,False
        elif a is not None and a!=0:amt,incoming=abs(a),a>0
        else:continue
        desc=get('description') or ''
        if not desc:
            # Preserve other descriptive columns, but don't copy amounts.
            used={v for k,v in cols.items() if k in ('date','amount','credit','debit') and v is not None}
            desc=' '.join(str(v) for j,v in enumerate(row) if j not in used and isinstance(v,str) and v.strip())
        out.append(proposal(dt,amt,desc,get('source') or '',incoming,line))
    return out

def read_delimited(text):
    sample=text[:6000]
    try:dialect=csv.Sniffer().sniff(sample,delimiters=',;\t|')
    except csv.Error:dialect=csv.excel
    return list(csv.reader(io.StringIO(text),dialect))

def pdf_transactions(content):
    import fitz
    doc=fitz.open(stream=content,filetype='pdf')
    if len(doc)>80:raise ValueError('PDF demasiado largo (máximo 80 páginas)')
    out=[]; previews=[]
    for page_num,page in enumerate(doc,1):
        text=page.get_text(sort=True);previews.append(text)
        words=page.get_text('words')
        # Infer each page's debit/credit/saldo positions from its column headings.
        headings={}
        for w in words:
            label=norm(w[4]).strip(':')
            if label in ('debito','credito','saldo','descripcion') and label not in headings:
                headings[label]=(w[0]+w[2])/2,w[1]
        if not all(h in headings for h in ('debito','credito','saldo')):continue
        xd=headings['debito'][0];xc=headings['credito'][0];xs=headings['saldo'][0]
        ystart=max(headings[k][1] for k in ('debito','credito','saldo'))+5
        lines={}
        for w in words:
            if w[1]<ystart:continue
            key=round((w[1]+w[3])/2 / 2)*2
            lines.setdefault(key,[]).append(w)
        year=2026;month=None
        hit=re.search(r'(?i)(enero|febrero|marzo|abril|mayo|junio|julio|agosto|septiembre|octubre|noviembre|diciembre)[/\s.-]*(\d{2,4})',text[:1600])
        months=['enero','febrero','marzo','abril','mayo','junio','julio','agosto','septiembre','octubre','noviembre','diciembre']
        if hit:month=months.index(hit[1].lower())+1;year=int(hit[2]);year=2000+year if year<100 else year
        # subsequent pages may lack a statement title; infer from the first page separately below
        for y,items in sorted(lines.items()):
            items.sort(key=lambda w:w[0])
            first=items[0][4] if items else ''
            if not re.fullmatch(r'\d{1,2}',first):continue
            dd=int(first)
            if not 1<=dd<=31:continue
            regions={'id':[],'description':[],'debit':[],'credit':[],'balance':[]}
            for w in items[1:]:
                x=(w[0]+w[2])/2
                if x < 79:regions['id'].append(w[4])
                elif x < xd-35:regions['description'].append(w[4])
                elif x < (xd+xc)/2:regions['debit'].append(w[4])
                elif x < (xc+xs)/2:regions['credit'].append(w[4])
                else:regions['balance'].append(w[4])
            debit=is_money(''.join(regions['debit'])) if regions['debit'] else None
            credit=is_money(''.join(regions['credit'])) if regions['credit'] else None
            if bool(debit) == bool(credit):continue
            # Month is resolved using the statement's first-page title below.
            desc=' '.join(regions['description'])
            if not desc:continue
            out.append({'day':dd,'page':page_num,'amount':credit or debit,'incoming':bool(credit),'desc':desc})
    full='\n'.join(previews)
    hit=re.search(r'(?i)(enero|febrero|marzo|abril|mayo|junio|julio|agosto|septiembre|octubre|noviembre|diciembre)\s*[/.-]\s*(\d{2,4})',full[:5000])
    if not hit:
        # Do not assign guessed dates to arbitrary statements.
        return [],full[:5000]
    month=months.index(hit[1].lower())+1;year=int(hit[2]);year=year+2000 if year<100 else year
    results=[]
    for i,r in enumerate(out):
        try:day=date(year,month,r['day']).isoformat()
        except ValueError:continue
        results.append(proposal(day,r['amount'],r['desc'],'',r['incoming'],f"p{r['page']}:{r['day']}:{i}"))
    return results,full[:5000]

def parse(file_name,content):
    suffix=Path(file_name).suffix.lower()
    if suffix not in ALLOWED:raise ValueError('Formato no admitido')
    if not content or len(content)>MAX_BYTES:raise ValueError('Archivo vacío o mayor a 12 MB')
    if suffix=='.csv':return structured_rows(read_delimited(content.decode('utf-8-sig',errors='replace'))),''
    if suffix=='.xlsx':
        from openpyxl import load_workbook
        wb=load_workbook(io.BytesIO(content),read_only=True,data_only=True)
        out=[]
        for ws in wb:
            rows=list(ws.iter_rows(values_only=True))[:3000]
            try:out.extend(structured_rows(rows))
            except ValueError:continue
        wb.close()
        if not out:raise ValueError('No encontré movimientos reconocibles en las hojas. Comprueba los datos o utiliza el registro manual.')
        return out,''
    if suffix=='.pdf':
        rows,preview=pdf_transactions(content)
        if rows:return rows,preview
        import fitz
        doc=fitz.open(stream=content,filetype='pdf')
        text='\n'.join(p.get_text(sort=True) for p in doc)
        if len(text.strip())<20:
            from PIL import Image
            text='\n'.join(ocr_image(Image.frombytes('RGB',(pix.width,pix.height),pix.samples)) for p in doc for pix in [p.get_pixmap(matrix=fitz.Matrix(1.5,1.5),colorspace=fitz.csRGB)])
    elif suffix=='.txt':text=content.decode('utf-8-sig',errors='replace')
    else:
        from PIL import Image
        text=ocr_image(Image.open(io.BytesIO(content)))
    return text_rows(text),text[:5000]

def text_rows(text):
    out=[]
    for i,line in enumerate(text.splitlines()[:3000]):
        m=DATE_RE.search(line)
        if not m:continue
        dt=is_date(m.group())
        if not dt:continue
        tail=line[m.end():]
        amounts=list(MONEY_RE.finditer(tail))
        if len(amounts)!=1:continue
        amt=is_money(amounts[0].group().strip())
        if not amt:continue
        desc=(tail[:amounts[0].start()]+tail[amounts[0].end():]).strip(' -|;:')
        out.append(proposal(dt,abs(amt),desc,'',amt>0,i+1))
    return out

def ocr_image(image):
    import pytesseract
    cmd=os.environ.get('TESSERACT_CMD')
    if cmd:pytesseract.pytesseract.tesseract_cmd=cmd
    try:return pytesseract.image_to_string(image,lang='spa+eng')
    except pytesseract.TesseractError:return pytesseract.image_to_string(image,lang='eng')
